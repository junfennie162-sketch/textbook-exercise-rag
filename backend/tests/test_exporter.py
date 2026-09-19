import docx as docx_lib

from app.services.exporter import export_solutions


def _sample_solutions() -> list[dict]:
    return [{
        "question_text": "计算 log2 8。",
        "answer_text": "## 解题思路\n利用对数定义。\n## 参考答案\n3",
        "sources": {
            "来源1": {"chunk_id": "c1", "source_file": "教材.pdf",
                      "chapter": "第三章", "page_number": 5,
                      "text_snippet": "对数的定义：…"},
        },
    }]


def test_export_markdown(tmp_path) -> None:
    path = export_solutions(_sample_solutions(), tmp_path, "md")
    content = path.read_text(encoding="utf-8")
    assert "习题解析报告" in content
    assert "第 1 题" in content
    assert "来源1" in content
    assert "第三章" in content


def test_export_docx_roundtrip(tmp_path) -> None:
    path = export_solutions(_sample_solutions(), tmp_path, "docx")
    doc = docx_lib.Document(str(path))
    texts = [p.text for p in doc.paragraphs]
    assert any("习题解析报告" in t for t in texts)
    assert any("第 1 题" in t for t in texts)
    assert any("第三章" in t for t in texts)


def test_export_includes_step_check(tmp_path) -> None:
    solutions = _sample_solutions()
    solutions[0]["step_check"] = {"checked": 2, "passed": 1, "skipped": 1,
                                  "failed": ["5 × 2 = 11"], "ok": False}

    md = export_solutions(solutions, tmp_path, "md").read_text(encoding="utf-8")
    assert "步骤验算：可验算 2 步，通过 1 步" in md
    assert "5 × 2 = 11" in md

    doc = docx_lib.Document(str(export_solutions(solutions, tmp_path, "docx")))
    assert any("步骤验算" in p.text for p in doc.paragraphs)


def test_export_skips_step_check_when_absent(tmp_path) -> None:
    md = export_solutions(_sample_solutions(), tmp_path, "md").read_text(encoding="utf-8")
    assert "步骤验算" not in md
