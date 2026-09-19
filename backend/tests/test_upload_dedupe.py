"""上传内容查重测试：同一份文件重复上传应被拦截（409），不同内容或不同资料类型正常入库。"""
import io

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.core.config import Settings
from app.main import app
from app.store import documents as doc_store
from app.store.db import init_db

client = TestClient(app)


class _FakeVectorStore:
    """隔离向量库：本测试只验证上传与查重逻辑，不加载嵌入模型。"""

    def __init__(self, settings=None):
        pass

    def add_chunks(self, chunks=None, **kwargs) -> int:
        return len(chunks or [])

    def delete_document(self, doc_id: str) -> int:
        return 0


class _FakeLexicalIndex:
    def add_chunks(self, chunks=None, **kwargs) -> int:
        return len(chunks or [])

    def remove_document(self, doc_id: str) -> int:
        return 0


@pytest.fixture
def isolated(tmp_path, monkeypatch):
    """数据目录指向临时目录，并隔离向量库与词法索引。

    注意：上传接口用的是 app.api.upload 内部的 get_settings，必须一并替换，
    否则上传文件会落到真实的 backend/data/uploads 下（数据库行在临时目录，文件成了孤儿）。
    """
    settings = Settings(data_dir=tmp_path)
    monkeypatch.setattr("app.store.db.get_settings", lambda: settings)
    monkeypatch.setattr("app.api.upload.get_settings", lambda: settings)

    from app.services import lexical_index as lex_mod
    from app.services import vector_store as vs_mod
    monkeypatch.setattr(vs_mod, "VectorStore", _FakeVectorStore)
    monkeypatch.setattr(lex_mod, "get_lexical_index", lambda: _FakeLexicalIndex())

    init_db()
    return settings


def _docx_bytes(text: str) -> bytes:
    buffer = io.BytesIO()
    document = Document()
    document.add_heading("第一章 集合与函数概念", level=1)
    document.add_paragraph(text)
    document.save(buffer)
    return buffer.getvalue()


def test_duplicate_upload_is_rejected(isolated) -> None:
    payload = _docx_bytes("集合具有确定性，例如不超过 20 的非负整数构成一个集合。")
    first = client.post("/api/upload/textbook", files={"file": ("教材.docx", payload)})
    assert first.status_code == 200
    assert first.json()["document"]["chunks_count"] >= 1

    # 同一内容换个文件名再传一次：应被拦截，且不新增文档
    second = client.post("/api/upload/textbook", files={"file": ("教材副本.docx", payload)})
    assert second.status_code == 409
    assert "内容已入库" in second.json()["detail"]
    assert "教材.docx" in second.json()["detail"]
    assert len(doc_store.load_documents()) == 1


def test_same_content_different_doc_type_allowed(isolated) -> None:
    """同一份文件既可作为教材、也可作为习题册登记，查重按「内容 + 类型」判定。"""
    payload = _docx_bytes("同一份内容在不同资料类型下分别登记。")
    assert client.post("/api/upload/textbook", files={"file": ("a.docx", payload)}).status_code == 200
    assert client.post("/api/upload/exercise", files={"file": ("a.docx", payload)}).status_code == 200
    assert len(doc_store.load_documents()) == 2


def test_different_content_allowed(isolated) -> None:
    assert client.post("/api/upload/textbook",
                       files={"file": ("a.docx", _docx_bytes("第一份不同的内容。"))}).status_code == 200
    assert client.post("/api/upload/textbook",
                       files={"file": ("b.docx", _docx_bytes("第二份不同的内容。"))}).status_code == 200
    assert len(doc_store.load_documents()) == 2


def test_document_record_keeps_content_hash(isolated) -> None:
    payload = _docx_bytes("用于校验内容指纹是否落库。")
    created = client.post("/api/upload/textbook", files={"file": ("c.docx", payload)}).json()["document"]
    stored = doc_store.get_document(created["doc_id"])
    assert stored["content_hash"] == created["content_hash"]
    assert len(stored["content_hash"]) == 64          # SHA-256 十六进制
    assert doc_store.find_document_by_hash(stored["content_hash"], "textbook")["doc_id"] == created["doc_id"]
    assert doc_store.find_document_by_hash(stored["content_hash"], "exercise") is None
