"""4.3 解析与引用来源导出：python-docx 生成 Word，Markdown 拼接。"""
import time
from pathlib import Path

import docx as docx_lib


def _step_check_line(sol: dict) -> str | None:
    """把步骤验算结果整理为一行说明（无可验算步骤时返回 None）。"""
    check = sol.get("step_check") or {}
    if not check.get("checked"):
        return None
    line = f"步骤验算：可验算 {check['checked']} 步，通过 {check['passed']} 步"
    if check.get("failed"):
        line += "；可疑步骤：" + "；".join(check["failed"])
    return line


def export_solutions(solutions: list[dict], out_dir: Path, fmt: str) -> Path:
    """把解析结果导出为 docx 或 md，返回文件路径。"""
    out_dir.mkdir(parents=True, exist_ok=True)
    stamp = time.strftime("%Y%m%d_%H%M%S")
    if fmt == "docx":
        return _export_docx(solutions, out_dir / f"solutions_export_{stamp}.docx")
    if fmt == "md":
        return _export_markdown(solutions, out_dir / f"solutions_export_{stamp}.md")
    raise ValueError(f"不支持的导出格式：{fmt}")


def _export_docx(solutions: list[dict], out_path: Path) -> Path:
    doc = docx_lib.Document()
    doc.add_heading("习题解析报告", level=0)
    doc.add_paragraph(f"共 {len(solutions)} 道题目，生成时间 {time.strftime('%Y-%m-%d %H:%M:%S')}")

    for idx, sol in enumerate(solutions, start=1):
        doc.add_heading(f"第 {idx} 题", level=1)
        doc.add_paragraph(sol.get("question_text", ""))
        doc.add_heading("解析", level=2)
        for line in (sol.get("answer_text") or "（无解析内容）").splitlines():
            if line.strip():
                doc.add_paragraph(line.rstrip())
        check_line = _step_check_line(sol)
        if check_line:
            doc.add_paragraph(check_line)
        doc.add_heading("引用来源", level=2)
        for key, ref in (sol.get("sources") or {}).items():
            doc.add_paragraph(
                f"{key}：《{ref.get('source_file', '')}》{ref.get('chapter', '')} "
                f"第{ref.get('page_number')}页 — {ref.get('text_snippet', '')}"
            )
    doc.save(str(out_path))
    return out_path


def _export_markdown(solutions: list[dict], out_path: Path) -> Path:
    lines = [f"# 习题解析报告\n", f"> 共 {len(solutions)} 道题目\n"]
    for idx, sol in enumerate(solutions, start=1):
        lines.append(f"## 第 {idx} 题\n")
        lines.append(f"**题目**：{sol.get('question_text', '')}\n")
        lines.append("### 解析\n")
        lines.append(f"{sol.get('answer_text') or '（无解析内容）'}\n")
        check_line = _step_check_line(sol)
        if check_line:
            lines.append(f"> {check_line}\n")
        lines.append("### 引用来源\n")
        for key, ref in (sol.get("sources") or {}).items():
            lines.append(
                f"- **{key}**：《{ref.get('source_file', '')}》{ref.get('chapter', '')} "
                f"第{ref.get('page_number')}页 — {ref.get('text_snippet', '')}"
            )
        lines.append("")
    out_path.write_text("\n".join(lines), encoding="utf-8")
    return out_path
